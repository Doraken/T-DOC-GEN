import tkinter as tk
from tkinter import simpledialog
from docx import Document

def creer_fichier_word(nom_app, code_app):
    # Créer un document Word
    doc = Document()
    doc.add_heading('Informations de l\'Application', 0)
    doc.add_paragraph(f'Nom de l\'Application: {nom_app}')
    doc.add_paragraph(f'Code du Projet: {code_app}')
    
    # Nommer le fichier en utilisant les entrées utilisateur
    nom_fichier = f'DAT_{nom_app}_{code_app}.docx'
    doc.save(nom_fichier)
    print(f'Fichier créé: {nom_fichier}')

def demander_informations():
    root = tk.Tk()
    root.withdraw()  # Cacher la fenêtre Tkinter principale
    
    # Demander le nom de l'application et le code du projet
    nom_app = simpledialog.askstring("Nom de l'Application", "Entrez le nom de l'application :")
    code_app = simpledialog.askstring("Code du Projet", "Entrez le code du projet :")
    
    root.destroy()  # Fermer la fenêtre Tkinter après la saisie des informations
    return nom_app, code_app

def run():
    nom_app, code_app = demander_informations()
    if nom_app and code_app:  # Vérifier que l'utilisateur a bien saisi les informations
        creer_fichier_word(nom_app, code_app)
    else:
        print("Opération annulée.")

if __name__ == "__main__":
    run()
